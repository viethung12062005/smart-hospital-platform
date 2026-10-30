"""System Analysis Result workbook: the sheet layout of the company template (List of requirement, Actor List,
Data Object Details, Metadata, Function List, Screen List), filled from the same SRS model so both stay in sync."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEAD = PatternFill("solid", fgColor="1F3A5F")
LABEL = PatternFill("solid", fgColor="DCE6F1")
THIN = Side(style="thin", color="9FB1C5")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
HFONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BFONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)

FUNCTION_TYPES = ("Basic: chức năng cơ bản (tạo, xem, sửa, xóa, tìm kiếm...) để người dùng làm việc với một đối tượng dữ "
                  "liệu.\nWorkflow: chức năng theo quy trình nghiệp vụ, ràng buộc chặt về trạng thái và vai trò (ví dụ: đặt "
                  "lịch, check-in, kết thúc lần khám).\nOther: chức năng khác (job định kỳ, AI, tích hợp).")


def _header(ws, row, col, names, widths):
    for i, (name, w) in enumerate(zip(names, widths)):
        c = ws.cell(row=row, column=col + i, value=name)
        c.fill, c.font, c.border, c.alignment = HEAD, HFONT, BOX, Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _row(ws, row, col, values, bold_first=False):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font = BOLD if (bold_first and i == 0) else BFONT
        c.border, c.alignment = BOX, WRAP


def write(m, path):
    r = m.resolve
    roles = {a["key"]: a["name"] for a in m["actors"]}
    wb = Workbook()

    ws = wb.active
    ws.title = "List of requirement"
    _header(ws, 2, 1, ["NO", "DESCRIPTION", "Ưu tiên", "Nguồn"], [6, 90, 10, 22])
    for i, req in enumerate(m["requirements"], 1):
        _row(ws, 2 + i, 1, [i, f"{req['id']} {req['title']}: {req['description']}", req["priority"], req["source"]], True)

    ws = wb.create_sheet("Actor List")
    _header(ws, 2, 2, ["NO", "Actor Name", "Description"], [6, 26, 120])
    row = 3
    for i, a in enumerate(m["actors"], 1):
        _row(ws, row, 2, [i, a["name"], a["description"]])
        row += 1
    row += 1
    _header(ws, row, 2, ["NO", "External System Name", "Description"], [6, 26, 120])
    for i, e in enumerate(m["external_systems"], 1):
        _row(ws, row + i, 2, [i, e["name"], e["description"]])

    ws = wb.create_sheet("Data Object Details")
    _header(ws, 2, 1, ["NO", "Actor", "Data Object", "Description", "Property", "Example", "Input Type", "Show on screen",
                       "Store on DB", "Notes"], [7, 22, 18, 40, 22, 22, 16, 10, 10, 34])
    row = 3
    for i, o in enumerate(m["objects"], 1):
        props = o.get("properties", [])
        first = row
        for p in props or [{}]:
            _row(ws, row, 1, [f"O{i:02d}" if row == first else None,
                              "\n".join(f"- {roles[a]}" for a in o.get("actors", [])) if row == first else None,
                              o["name"] if row == first else None, o["description"] if row == first else None,
                              p.get("name"), p.get("example"), p.get("input"), "x" if p.get("show", True) else "",
                              "x" if p.get("store", True) else "", p.get("note")])
            row += 1
        if len(props) > 1:
            for col in range(1, 5):
                ws.merge_cells(start_row=first, start_column=col, end_row=row - 1, end_column=col)

    ws = wb.create_sheet("Metadata")
    ws["A1"] = "Bảng mô tả loại của từng thuộc tính trong sheet Data Object Details."
    _header(ws, 2, 1, ["Type", "Description"], [20, 90])
    for i, (t, d) in enumerate((("Free Input", "Người dùng nhập tự do theo định dạng quy định (văn bản, số, ngày...)."),
                                ("Input in a list", "Người dùng chỉ chọn từ danh sách giá trị định sẵn (value set)."),
                                ("System generated", "Hệ thống tự sinh hoặc tự cập nhật (mã, thời điểm, trạng thái...)."),
                                ("From external system", "Đồng bộ từ hệ thống bên ngoài (nhà cung cấp định danh...).")), 1):
        _row(ws, 2 + i, 1, [t, d])

    ws = wb.create_sheet("Function List")
    fn_rows = []
    n = 0
    for module, section, uc in m.use_cases():
        n += 1
        fn_rows.append([f"F{n:02d}", module["name"], r(uc.get("screen_text") or uc["trigger"]),
                        f"{m.num[('UC', uc['key'])]}: {uc['name']}", uc.get("type", "Workflow"), r(uc["objective"])])
    ws["B2"], ws["C2"] = "Total Number of Function", len(fn_rows)
    ws["B2"].font, ws["C2"].font = BOLD, BOLD
    ws["G2"] = FUNCTION_TYPES
    ws["G2"].alignment = WRAP
    ws.row_dimensions[2].height = 60
    _header(ws, 4, 2, ["Trace Code", "Category (module)", "Trigger / Screen", "Function", "Type", "Description / Objective / Remarks"],
            [11, 22, 40, 40, 12, 70])
    for i, fr in enumerate(fn_rows, 1):
        _row(ws, 4 + i, 2, fr, True)

    ws = wb.create_sheet("Screen List")
    sc_rows = []
    for portal in m["screens"]:
        for s in portal["items"]:
            sc_rows.append([m.num[("SCR", s["key"])].replace("Màn hình ", "S"), portal["portal"], s["name"], r(s["purpose"])])
    ws["B2"], ws["C2"] = "Total Number of screen", len(sc_rows)
    ws["B2"].font, ws["C2"].font = BOLD, BOLD
    _header(ws, 4, 2, ["Trace Code", "Category (module)", "Screen", "Description / Objective / Remarks"], [11, 24, 36, 90])
    for i, sr in enumerate(sc_rows, 1):
        _row(ws, 4 + i, 2, sr, True)

    wb.save(path)

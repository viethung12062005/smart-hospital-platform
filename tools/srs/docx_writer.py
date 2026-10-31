"""Render the SRS model as a .docx that follows the section structure of the company SRS template
(cover, approval page, revision history, table of contents, chapters 1-9). No company branding is used."""
from __future__ import annotations

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x59, 0x59, 0x59)
HEAD_FILL = "1F3A5F"
LABEL_FILL = "DCE6F1"
FONT = "Arial"


class Writer:
    def __init__(self, model, pngs):
        self.m = model
        self.pngs = pngs
        self.doc = Document()
        self.fig = 0
        self.h = [0, 0, 0, 0]
        self._styles()

    # ------------------------------------------------------------------ styles and helpers
    def _styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name, normal.font.size = FONT, Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.15
        for lvl, size in ((1, 16), (2, 13), (3, 11.5), (4, 11)):
            s = st[f"Heading {lvl}"]
            s.font.name, s.font.size, s.font.bold, s.font.color.rgb = FONT, Pt(size), True, NAVY
            s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            s.paragraph_format.space_before = Pt(14 if lvl <= 2 else 10)
            s.paragraph_format.space_after = Pt(6)
            s.paragraph_format.keep_with_next = True
        cap = st["Caption"]
        cap.font.name, cap.font.size, cap.font.italic, cap.font.bold, cap.font.color.rgb = FONT, Pt(9.5), True, False, GREY
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2)

    def r(self, text):
        return self.m.resolve(text)

    def heading(self, text, level):
        self.h[level - 1] += 1
        for i in range(level, 4):
            self.h[i] = 0
        num = ".".join(str(x) for x in self.h[:level])
        return self.doc.add_heading(f"{num}  {text}", level)

    def para(self, text="", bold=False, italic=False, size=None, color=None, align=None, style=None):
        p = self.doc.add_paragraph(style=style)
        if text:
            run = p.add_run(self.r(text))
            run.bold, run.italic = bold, italic
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color
        if align:
            p.alignment = align
        return p

    def bullets(self, items, level=0):
        for it in items or []:
            if isinstance(it, dict):
                self.bullets([it["text"]], level)
                self.bullets(it.get("items"), level + 1)
                continue
            p = self.doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
            p.add_run(self.r(it))

    def note(self, text):
        p = self.para(text, italic=True, color=GREY, size=9.5)
        return p

    def figure(self, key, caption, width_cm=16.0):
        path = self.pngs.get(key)
        self.fig += 1
        if path and os.path.exists(path):
            from PIL import Image
            with Image.open(path) as im:
                w, h = im.size
            width = min(width_cm, 16.5)
            if h / w * width > 22:                      # keep very tall diagrams on one page
                width = 22 * w / h
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(path, width=Cm(width))
        else:
            self.note(f"[Hình chưa được sinh: {key}]")
        c = self.doc.add_paragraph(f"Hình {self.fig}: {caption}", style="Caption")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER

    @staticmethod
    def shade(cell, fill):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcpr.append(shd)

    def cell_text(self, cell, value, bold=False, color=None, size=9.5):
        cell.text = ""
        lines = value if isinstance(value, list) else [value]
        first = True
        for line in lines:
            if isinstance(line, dict):          # {title, lines} block inside a cell
                self.cell_text_append(cell, line["title"], bold=True, first=first, size=size)
                first = False
                for sub in line.get("lines", []):
                    self.cell_text_append(cell, f"• {sub}", first=False, size=size)
                continue
            self.cell_text_append(cell, str(line), bold=bold, color=color, first=first, size=size)
            first = False

    def cell_text_append(self, cell, text, bold=False, color=None, first=False, size=9.5):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(self.r(text))
        run.bold = bold
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color

    def table(self, header, rows, widths=None, label_col=False, font=9.5):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            self.cell_text(c, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=font)
            self.shade(c, HEAD_FILL)
        self._repeat_header(t.rows[0])
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                self.cell_text(cells[i], v if v is not None else "", bold=(label_col and i == 0), size=font)
                if label_col and i == 0:
                    self.shade(cells[i], LABEL_FILL)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def kv_table(self, rows, widths=(4.0, 12.5)):
        t = self.doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, v in rows:
            cells = t.add_row().cells
            self.cell_text(cells[0], k, bold=True)
            self.shade(cells[0], LABEL_FILL)
            self.cell_text(cells[1], v)
        for row in t.rows:
            row.cells[0].width, row.cells[1].width = Cm(widths[0]), Cm(widths[1])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    @staticmethod
    def _repeat_header(row):
        trpr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trpr.append(el)

    @staticmethod
    def field(paragraph, instr):
        run = paragraph.add_run()
        for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "separate"), ("w:t", None),
                          ("w:fldChar", "end")):
            el = OxmlElement(tag)
            if tag == "w:fldChar":
                el.set(qn("w:fldCharType"), attr)
            elif tag == "w:instrText":
                el.set(qn("xml:space"), "preserve")
                el.text = instr
            else:
                el.text = "1"
            run._r.append(el)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------ front matter
    def front(self):
        d = self.m["document"]
        for _ in range(6):
            self.para()
        self.para(d["system"].upper(), bold=True, size=24, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.para("ĐẶC TẢ YÊU CẦU PHẦN MỀM", bold=True, size=18, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.para("(Software Requirements Specification)", italic=True, size=12, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.para()
        self.para(d["subtitle"], size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        for _ in range(8):
            self.para()
        self.para(f"Phiên bản: {d['version']}", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.para(d["team"], align=WD_ALIGN_PARAGRAPH.CENTER)
        self.para(d["date_text"], align=WD_ALIGN_PARAGRAPH.CENTER)
        self.page_break()

        self.para("Trang phê duyệt", bold=True, size=14, color=NAVY)
        self.para(d["approval_text"])
        t = self.doc.add_table(rows=0, cols=4)
        t.style = "Table Grid"
        for role, who in d["approvals"]:
            for i, (a, b, c_, e) in enumerate(((f"{role}:", who, "Chữ ký:", "____________________"),
                                               ("", "", "Ngày:", "____ / ____ / ________"))):
                cells = t.add_row().cells
                for j, v in enumerate((a, b, c_, e)):
                    self.cell_text(cells[j], v, bold=(j == 0))
        self.para()
        self.para("Lịch sử thay đổi", bold=True, size=14, color=NAVY)
        self.table(["Ngày", "Phiên bản", "Người thực hiện", "Mô tả thay đổi"],
                   [[h["date"], h["version"], h["author"], h["change"]] for h in d["history"]], widths=(2.6, 2.0, 3.6, 8.3))
        self.page_break()
        self.para("Mục lục", bold=True, size=14, color=NAVY)
        p = self.doc.add_paragraph()
        self.field(p, 'TOC \\o "1-4" \\h \\z \\u')
        self.note("Nhấn chuột phải vào mục lục → Update Field nếu số trang chưa đúng.")
        self.page_break()

    def header_footer(self):
        d = self.m["document"]
        sec = self.doc.sections[0]
        hp = sec.header.paragraphs[0]
        hp.text = f"{d['team']} – {d['system']}\tĐặc tả yêu cầu phần mềm"
        for run in hp.runs:
            run.font.size, run.font.color.rgb = Pt(8.5), GREY
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        run = fp.add_run(f"Người soạn: {d['prepared_by']}   |   Phiên bản {d['version']}   |   Trang ")
        run.font.size, run.font.color.rgb = Pt(8.5), GREY
        self.field(fp, "PAGE")
        run = fp.add_run(" / ")
        run.font.size = Pt(8.5)
        self.field(fp, "NUMPAGES")
        settings = self.doc.settings.element
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)

    # ------------------------------------------------------------------ chapter 1
    def introduction(self):
        intro = self.m["introduction"]
        self.heading("Giới thiệu", 1)
        self.heading("Mục đích", 2)
        self.para(intro["purpose"])
        self.para("Cụ thể, tài liệu này:")
        self.bullets(intro["purpose_points"])
        self.heading("Tổng quan", 2)
        for p in intro["overview"]:
            self.para(p)
        self.heading("Đối tượng đọc và cách đọc", 2)
        self.para("Tài liệu dành cho:")
        self.bullets(intro["audience"])
        self.para(intro["reading"])
        self.heading("Thuật ngữ và viết tắt", 2)
        self.table(["Viết tắt / Quy ước", "Ý nghĩa"], sorted(intro["glossary"], key=lambda x: x[0].lower()), widths=(4.2, 12.3))
        self.heading("Tài liệu tham khảo", 2)
        self.table(["Tài liệu", "Vị trí", "Mô tả"], intro["references"], widths=(4.5, 5.5, 6.5))

    # ------------------------------------------------------------------ chapter 2
    def high_level(self):
        hl = self.m["high_level"]
        self.heading("Yêu cầu mức tổng quan", 1)
        self.para(hl["intro"])
        self.heading("Sơ đồ quan hệ đối tượng", 2)
        self.para(hl["ord"]["intro"])
        self.figure("ord", "Sơ đồ quan hệ đối tượng")
        self.para("Mô tả đối tượng:", bold=True)
        rows = []
        for group, items in (("Đối tượng dữ liệu", self.m["objects"]), ("Actor", self.m["actors"]),
                             ("Hệ thống bên ngoài", self.m["external_systems"])):
            rows.append([group, "", ""])
            for i, o in enumerate(items, 1):
                rows.append([str(i), o["name"], o["description"]])
        t = self.table(["#", "Đối tượng", "Mô tả"], rows, widths=(1.0, 3.8, 11.7))
        for row in t.rows[1:]:
            if row.cells[1].text == "" and row.cells[0].text:
                merged = row.cells[0].merge(row.cells[2])
                self.shade(merged, LABEL_FILL)
                for p in merged.paragraphs:
                    for run in p.runs:
                        run.bold = True
        self.heading("Quy trình nghiệp vụ", 2)
        self.para(hl["workflow_intro"])
        for wf in hl["workflows"]:
            self.heading(wf["name"], 3)
            self.figure(f"wf-{wf['key']}", wf["name"])
            self.para("Giải thích quy trình:", bold=True)
            self.bullets(wf["explanation"])
        self.heading("Chuyển trạng thái", 2)
        self.para(hl["state_intro"])
        for st in hl["state_machines"]:
            self.heading(st["name"], 3)
            self.figure(f"st-{st['key']}", st["name"])
            self.table(["Trạng thái", "Ý nghĩa"], [[s["label"], s["meaning"]] for s in st["states"]], widths=(4, 12.5))
        self.heading("Sơ đồ use case", 2)
        self.para(hl["uc_intro"])
        for g in self.m["uc_diagrams"]:
            self.heading(g["name"], 3)
            self.figure(f"ucd-{g['key']}", g["name"])
            uc = {u["key"]: u for _, _, u in self.m.use_cases()}
            self.table(["#", "Use case", "Mô tả"],
                       [[self.m.num[("UC", k)], uc[k]["name"], uc[k]["objective"]] for k in g["cases"]],
                       widths=(1.6, 4.6, 10.3))
        self.permission_matrix()
        self.heading("Sơ đồ trang (Site map)", 2)
        for sm in hl["site_maps"]:
            self.heading(sm["name"], 3)
            self.para("Sơ đồ trang mô tả cách điều hướng trong hệ thống (tối đa 3 cấp tính từ trang chủ).")
            self.figure(f"sm-{sm['key']}", sm["name"])
            self.table(["Trang", "Mô tả", "Quyền truy cập"], sm["pages"], widths=(3.8, 8.7, 4.0))

    def permission_matrix(self):
        pm = self.m["permissions"]
        self.heading("Ma trận phân quyền", 2)
        self.para(pm["intro"])
        self.para("Ký hiệu:", bold=True)
        self.bullets(pm["legend"])
        roles = [r for r in self.m["actors"] if r.get("in_matrix", True)]
        rows = [[row["function"]] + [row["grant"].get(r["key"], "X") for r in roles] for row in pm["rows"]]
        t = self.table(["Chức năng"] + [r["short"] for r in roles], rows,
                       widths=[6.0] + [10.5 / len(roles)] * len(roles), font=9)
        for row in t.rows[1:]:
            for c in row.cells[1:]:
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.para("Ghi chú:", bold=True)
        self.bullets(pm["notes"])

    # ------------------------------------------------------------------ chapter 3
    def use_cases(self):
        self.heading("Đặc tả use case", 1)
        self.para(self.m["use_case_intro"])
        for module in self.m["modules"]:
            self.heading(module["name"], 2)
            if module.get("intro"):
                self.para(module["intro"])
            for section in module["sections"]:
                self.heading(section["name"], 3)
                for uc in section["use_cases"]:
                    self.use_case(uc)
        self.common_use_cases()
        self.heading("Quy tắc nghiệp vụ chung", 2)
        self.para("Các quy tắc dưới đây dùng chung cho nhiều use case; mỗi quy tắc chỉ áp dụng cho use case có tham chiếu tới nó.")
        self.table(["Mã", "Mô tả"], [[self.m.num[("CBR", r["key"])], [{"title": r["title"], "lines": r["lines"]}]]
                                     for r in self.m["common_rules"]], widths=(2.0, 14.5))

    def use_case(self, uc):
        label = self.m.num[("UC", uc["key"])]
        self.heading(f"{label}: {uc['name']}", 4)
        roles = {r["key"]: r["name"] for r in self.m["actors"]}
        self.kv_table([
            ("Mục tiêu:", uc["objective"]),
            ("Actor:", ", ".join(roles[a] for a in uc["actors"]) + "."),
            ("Kích hoạt:", uc["trigger"]),
            ("Điều kiện trước:", uc["pre"]),
            ("Điều kiện sau:", uc["post"]),
            ("Truy vết:", ", ".join(uc["trace"])),
        ])
        if uc.get("steps"):
            self.para("Luồng hoạt động", bold=True)
            self.figure(f"act-{uc['key']}", f"Luồng hoạt động – {label}: {uc['name']}")
        if uc.get("alternatives"):
            self.para("Luồng thay thế và ngoại lệ", bold=True)
            self.bullets(uc["alternatives"])
        if uc.get("rules"):
            self.para("Quy tắc nghiệp vụ", bold=True)
            self.table(["Bước", "Mã BR", "Mô tả"],
                       [[br["step"], br["_code"], [{"title": br["title"], "lines": br["lines"]}]] for br in uc["rules"]],
                       widths=(1.8, 1.7, 13.0))
        if uc.get("acceptance"):
            self.para("Tiêu chí nghiệm thu", bold=True)
            self.table(["Mã", "Tiêu chí (Given / When / Then)"], uc["acceptance"], widths=(2.2, 14.3))

    def common_use_cases(self):
        self.heading("Use case dùng chung", 2)
        self.para("Các chức năng chuẩn khi làm việc với một đối tượng dữ liệu. Mỗi use case dùng chung mô tả quy tắc điển "
                  "hình; use case cụ thể chỉ ghi phần khác biệt.")
        for c in self.m["common_use_cases"]:
            label = self.m.num[("CMUC", c["key"])]
            self.heading(f"{label}: {c['name']}", 3)
            self.kv_table([("Mục tiêu:", c["objective"]), ("Actor:", c["actor"]), ("Kích hoạt:", c["trigger"]),
                           ("Điều kiện trước:", c["pre"]), ("Điều kiện sau:", c["post"])])
            if c.get("steps"):
                self.para("Luồng hoạt động", bold=True)
                self.figure(f"act-{c['key']}", f"Luồng hoạt động – {label}: {c['name']}")
            if c.get("rules"):
                self.para("Quy tắc nghiệp vụ", bold=True)
                self.table(["Bước", "Mã BR", "Mô tả"],
                           [[br["step"], br["_code"], [{"title": br["title"], "lines": br["lines"]}]] for br in c["rules"]],
                           widths=(1.8, 1.7, 13.0))

    # ------------------------------------------------------------------ chapter 4
    def screens(self):
        self.heading("Màn hình (Mockup)", 1)
        self.para(self.m["screens_intro"])
        for portal in self.m["screens"]:
            self.heading(portal["portal"], 2)
            for s in portal["items"]:
                label = self.m.num[("SCR", s["key"])]
                self.heading(s["name"], 3)
                self.bullets([s["purpose"]])
                self.bullets(["Truy cập: " + s["access"]])
                if s.get("uc"):
                    self.bullets(["Use case liên quan: " + ", ".join(f"[[UC:{k}]]" for k in s["uc"])])
                key = f"mock-{s['key']}"
                if key in self.pngs:
                    self.figure(key, f"{label}: {s['name']}")
                else:
                    self.note(f"{label}: {s['name']} – mockup bổ sung ở phiên bản sau.")
                if s.get("data_source") or s.get("sorting"):
                    self.kv_table([(k, v) for k, v in (("Nguồn dữ liệu", s.get("data_source")),
                                                      ("Sắp xếp mặc định", s.get("sorting"))) if v])
                if s.get("components"):
                    self.table(["#", "Thành phần", "Loại", "Sửa được", "Bắt buộc", "Mặc định", "Mô tả"],
                               [[str(i)] + c for i, c in enumerate(s["components"], 1)],
                               widths=(0.8, 3.0, 2.0, 1.5, 1.5, 2.0, 5.7), font=8.5)

    # ------------------------------------------------------------------ chapter 5-9
    def nfr(self):
        n = self.m["nfr"]
        self.heading("Yêu cầu phi chức năng", 1)
        self.para(n["intro"])
        for cat in n["categories"]:
            self.heading(cat["name"], 2)
            if cat.get("intro"):
                self.para(cat["intro"])
            self.table(["Mã", "Yêu cầu", "Chỉ số / ngưỡng", "Cách đo và điều kiện", "Kiểm thử"],
                       [[x["id"], x["title"], x["target"], x["measure"], x.get("test", "")] for x in cat["items"]],
                       widths=(1.7, 3.6, 4.2, 4.8, 2.2), font=8.5)

    def other(self):
        o = self.m["other"]
        self.heading("Yêu cầu khác", 1)
        for sec in o:
            self.heading(sec["name"], 2)
            if sec.get("intro"):
                self.para(sec["intro"])
            if sec.get("table"):
                self.table(sec["table"]["header"], sec["table"]["rows"], widths=sec["table"].get("widths"))
            self.bullets(sec.get("bullets"))

    def integration(self):
        self.heading("Tích hợp", 1)
        self.para(self.m["integration"]["intro"])
        self.table(["Hệ thống", "Chiều tích hợp", "Giao thức / dữ liệu", "Xử lý khi lỗi"], self.m["integration"]["rows"],
                   widths=(3.5, 2.6, 5.2, 5.2))

    def migration(self):
        mg = self.m["migration"]
        self.heading("Chuyển đổi dữ liệu", 1)
        self.para(mg["intro"])
        self.heading("Phạm vi", 2)
        self.bullets(mg["scope"])
        self.heading("Ánh xạ dữ liệu", 2)
        self.para(mg["mapping"])

    def appendices(self):
        self.heading("Phụ lục", 1)
        self.heading("Danh sách thông điệp", 2)
        for group in self.m["messages"]:
            self.heading(group["name"], 3)
            self.para(group["intro"])
            self.bullets(group.get("format"))
            self.table(["Mã", "Nội dung", "Khi nào hiển thị"],
                       [[self.m.num[("MSG", x["key"])], x["text"], x.get("when", "")] for x in group["items"]],
                       widths=(1.8, 8.2, 6.5), font=9)
        self.heading("Mẫu thông báo (email / SMS / push)", 2)
        self.para(self.m["email_intro"])
        for et in self.m["email_templates"]:
            self.heading(f"{self.m.num[('ET', et['key'])]}: {et['name']}", 3)
            self.kv_table([("Kênh", et["channel"]), ("Người nhận", et["to"]), ("Kích hoạt", et["trigger"]),
                           ("Tiêu đề", et.get("subject", "N/A")), ("Nội dung", et["body"])])
            if et.get("notes"):
                self.para("Ghi chú:", bold=True)
                self.bullets(et["notes"])

    def rationale(self):
        r = self.m["rationale"]
        self.heading("Cơ sở đặt ngưỡng và căn cứ pháp lý", 2)
        self.para(r["intro"])
        self.para("Mức chắc chắn:", bold=True)
        self.bullets(r["confidence_legend"])
        self.heading("Các ngưỡng và giá trị", 3)
        self.table(["Tham số", "Giá trị đã chọn", "Căn cứ và lý do", "Mức chắc chắn"], r["thresholds"],
                   widths=(3.0, 3.6, 7.4, 2.5), font=8.5)
        self.heading("Văn bản pháp luật liên quan", 3)
        self.table(["Văn bản", "Nội dung liên quan", "Điều khoản / trạng thái"], r["legal"], widths=(4.2, 8.0, 4.3), font=9)
        self.heading("Việc còn phải xác minh", 3)
        self.bullets(r["open_items"])
        self.heading("Nguồn tham khảo (tra cứu 31/10/2026)", 3)
        self.table(["Nguồn", "Liên kết"], r["sources"], widths=(7.0, 9.5), font=8.5)

    def build(self, path):
        self.header_footer()
        self.front()
        self.introduction()
        self.high_level()
        self.use_cases()
        self.screens()
        self.nfr()
        self.other()
        self.integration()
        self.migration()
        self.appendices()
        self.rationale()
        self.doc.core_properties.title = f"{self.m['document']['system']} - Đặc tả yêu cầu phần mềm"
        self.doc.core_properties.author = self.m["document"]["team"]
        self.doc.save(path)


def write(model, pngs, path):
    Writer(model, pngs).build(path)

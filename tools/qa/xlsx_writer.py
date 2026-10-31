"""Test case workbook for testers and for tracking automation progress."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEAD = PatternFill("solid", fgColor="1F3A5F")
THIN = Side(style="thin", color="C9D3DD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS_FILL = {"automated": "C6EFCE", "partial": "FFE699", "planned": "F8CBAD", "manual": "DDEBF7"}


def _header(ws, names, widths):
    for i, (n, w) in enumerate(zip(names, widths), 1):
        c = ws.cell(row=1, column=i, value=n)
        c.fill, c.font, c.border = HEAD, Font(bold=True, color="FFFFFF", name="Arial", size=10), BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "C2"


def write_cases(path, rows, by_req):
    wb = Workbook()
    ws = wb.active
    ws.title = "Test cases"
    cols = ["ID", "Tên test case", "Yêu cầu phủ", "Module", "Mức", "Ưu tiên", "Tự động hóa", "Test tự động", "Thiếu gì / chờ gì", "Người phụ trách", "Mô tả (Given/When/Then)"]
    _header(ws, cols, [20, 44, 18, 14, 12, 9, 13, 50, 44, 18, 80])
    for i, r in enumerate(sorted(rows, key=lambda x: x["id"]), 2):
        refs = [x.split("/")[-1] for x in r["refs"]] + [f"{t['file']}::{t['title'][:60]}" for t in r["e2e"]]
        vals = [r["id"], r["title"], ", ".join(r["covers"]), r["module"], r["level"], r["priority"], r["status"], "\n".join(refs),
                r["gap"], r["owner"], r["spec"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font, cell.border, cell.alignment = Font(name="Arial", size=9), BOX, WRAP
    ws.auto_filter.ref = f"A1:K{ws.max_row}"
    for status, color in STATUS_FILL.items():
        ws.conditional_formatting.add(f"G2:G{ws.max_row}", CellIsRule(operator="equal", formula=[f'"{status}"'], fill=PatternFill("solid", fgColor=color)))

    ws2 = wb.create_sheet("Truy vết")
    _header(ws2, ["Yêu cầu (AC / NFR)", "Test case", "Mức", "Ưu tiên", "Tự động hóa"], [22, 28, 14, 9, 14])
    ws2.freeze_panes = "A2"
    row = 2
    for req in sorted(by_req):
        for r in by_req[req]:
            for c, v in enumerate([req, r["id"], r["level"], r["priority"], r["status"]], 1):
                cell = ws2.cell(row=row, column=c, value=v)
                cell.font, cell.border = Font(name="Arial", size=9), BOX
            row += 1
    ws2.auto_filter.ref = f"A1:E{ws2.max_row}"

    ws3 = wb.create_sheet("Tổng hợp")
    ws3["A1"] = "Tổng hợp theo trạng thái tự động hóa (công thức đếm từ sheet 'Test cases')"
    ws3["A1"].font = Font(bold=True, name="Arial", size=11)
    heads = ["Trạng thái", "Số test case", "Trong đó P0"]
    for i, h in enumerate(heads, 1):
        c = ws3.cell(row=3, column=i, value=h)
        c.fill, c.font, c.border = HEAD, Font(bold=True, color="FFFFFF", name="Arial", size=10), BOX
    for i, st in enumerate(["automated", "partial", "planned", "manual"], 4):
        ws3.cell(row=i, column=1, value=st).border = BOX
        ws3.cell(row=i, column=2, value=f"=COUNTIF('Test cases'!$G:$G,A{i})").border = BOX
        ws3.cell(row=i, column=3, value=f"=COUNTIFS('Test cases'!$G:$G,A{i},'Test cases'!$F:$F,\"P0\")").border = BOX
    ws3.cell(row=8, column=1, value="Tổng").font = Font(bold=True)
    ws3.cell(row=8, column=2, value="=SUM(B4:B7)")
    ws3.cell(row=8, column=3, value="=SUM(C4:C7)")
    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["B"].width = 16
    ws3.column_dimensions["C"].width = 14
    wb.save(path)
